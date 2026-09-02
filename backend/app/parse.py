"""Parse uploaded tournament data (Excel .xlsx or CSV) into a normalized form.

Two logical tables are expected (as two sheets in one .xlsx, or two CSV files):

PLAYERS  — one row per player
    start_no   (int, required)  the Swiss Manager starting rank / seed
    name       (required)       "Lastname, Firstname"
    rating     (int, optional)
    fide_id    (optional)
    sex        (m/w, optional)
    title      (optional)
    federation (optional)

PAIRINGS — one row per board per round (the ACTUAL Swiss Manager pairings)
    round   (int, required)
    board   (int, optional)     display order only
    white   (int, required)     start_no of the white player
    black   (int, optional)     start_no of the black player; empty/0 = bye
    result  (required)          white-perspective: 1-0 | 0.5-0.5 | 0-1
                                forfeits: 1-0F | 0-1F ; bye rows: bye | U | H | F | Z

Column names are matched case-insensitively with common aliases, so exports
don't have to be renamed perfectly.
"""

from __future__ import annotations

import csv
import io
import json
from typing import Optional

from openpyxl import load_workbook

# ---- column aliases (all lowercased, non-alnum stripped) ----
_PLAYER_ALIASES = {
    "start_no": {"startno", "start", "seed", "sno", "no", "rank", "startingrank", "startrank", "id", "nr"},
    "name": {"name", "playername", "player", "fullname"},
    "rating": {"rating", "elo", "rtg", "fiderating"},
    "fide_id": {"fideid", "fide", "fideno", "idfide", "fidenumber"},
    "sex": {"sex", "gender", "s"},
    "title": {"title", "ttl"},
    "federation": {"federation", "fed", "country", "nation"},
}
_PAIRING_ALIASES = {
    "round": {"round", "rnd", "r", "roundno", "roundnumber"},
    "board": {"board", "bo", "boardno", "table", "tableno"},
    "white": {"white", "whiteno", "whitestartno", "w", "whiteseed", "whiteid"},
    "black": {"black", "blackno", "blackstartno", "b", "blackseed", "blackid"},
    "result": {"result", "res", "score", "outcome", "whiteresult"},
}


def _norm(s: str) -> str:
    return "".join(ch for ch in str(s).strip().lower() if ch.isalnum())


def _resolve(headers: list[str], aliases: dict) -> dict:
    """Map our canonical field -> the actual column index in this file."""
    normed = [_norm(h) for h in headers]
    out = {}
    for field, names in aliases.items():
        for i, h in enumerate(normed):
            if h == field.replace("_", "") or h in names:
                out[field] = i
                break
    return out


def _rows_from_xlsx_sheet(ws) -> list[list]:
    return [[c.value for c in row] for row in ws.iter_rows()]


def _rows_from_csv(text: str) -> list[list]:
    return [row for row in csv.reader(io.StringIO(text)) if any(str(c).strip() for c in row)]


def _to_int(v) -> Optional[int]:
    if v is None or str(v).strip() == "":
        return None
    try:
        return int(float(str(v).strip()))
    except ValueError:
        return None


def parse_players(rows: list[list]) -> list[dict]:
    if not rows:
        raise ValueError("players table is empty")
    idx = _resolve(rows[0], _PLAYER_ALIASES)
    for req in ("start_no", "name"):
        if req not in idx:
            raise ValueError(f"players table missing a '{req}' column (got {rows[0]})")
    players = []
    for r in rows[1:]:
        if idx["start_no"] >= len(r):
            continue
        sno = _to_int(r[idx["start_no"]])
        name = r[idx["name"]] if idx["name"] < len(r) else None
        if sno is None or not name:
            continue
        players.append({
            "start_no": sno,
            "name": str(name).strip(),
            "rating": _to_int(r[idx["rating"]]) if "rating" in idx and idx["rating"] < len(r) else 0,
            "fide_id": (str(r[idx["fide_id"]]).strip() if "fide_id" in idx and idx["fide_id"] < len(r) and r[idx["fide_id"]] not in (None, "") else None),
            "sex": (str(r[idx["sex"]]).strip().lower()[:1] if "sex" in idx and idx["sex"] < len(r) and r[idx["sex"]] else None),
            "title": (str(r[idx["title"]]).strip() if "title" in idx and idx["title"] < len(r) and r[idx["title"]] else None),
            "federation": (str(r[idx["federation"]]).strip() if "federation" in idx and idx["federation"] < len(r) and r[idx["federation"]] else None),
        })
    if not players:
        raise ValueError("no player rows parsed")
    return players


def parse_pairings(rows: list[list]) -> list[dict]:
    if not rows:
        raise ValueError("pairings table is empty")
    idx = _resolve(rows[0], _PAIRING_ALIASES)
    for req in ("round", "white", "result"):
        if req not in idx:
            raise ValueError(f"pairings table missing a '{req}' column (got {rows[0]})")
    out = []
    for r in rows[1:]:
        if idx["round"] >= len(r):
            continue
        rnd = _to_int(r[idx["round"]])
        white = _to_int(r[idx["white"]]) if idx["white"] < len(r) else None
        if rnd is None or white is None:
            continue
        black = _to_int(r[idx["black"]]) if "black" in idx and idx["black"] < len(r) else None
        if black == 0:
            black = None
        result = str(r[idx["result"]]).strip() if idx["result"] < len(r) and r[idx["result"]] is not None else ""
        board = _to_int(r[idx["board"]]) if "board" in idx and idx["board"] < len(r) else None
        out.append({"round": rnd, "board": board, "white": white, "black": black, "result": result})
    if not out:
        raise ValueError("no pairing rows parsed")
    return out


# --------------------------- Chess-Results.com format ---------------------- #
def _is_chess_results(rows: list[list]) -> bool:
    """A Chess-Results.com export: identified by its site header text, or by a
    standalone 'Round N' label row combined with a White/Black table header.
    (Our own two-sheet Players/Pairings template must NOT match — it has a
    White/Black header but no 'Round N' label row and no site text.)"""
    head = " ".join(str(c) for r in rows[:25] for c in r if c is not None).lower()
    if "chess-results" in head:
        return True
    has_round_marker = any(
        r and r[0] is not None and str(r[0]).strip().lower().startswith("round")
        and any(ch.isdigit() for ch in str(r[0]))
        for r in rows[:300]
    )
    has_wb_header = any(
        "white" in [str(c).strip().lower() if c is not None else "" for c in r]
        and "black" in [str(c).strip().lower() if c is not None else "" for c in r]
        for r in rows[:40]
    )
    return has_round_marker and has_wb_header


def _cr_result(raw: str, has_black: bool):
    """Map a Chess-Results result cell to a white-perspective code, or a bye code.

    Byes/not-paired (has_black=False) return a bye letter U/H/F/Z. Chess-Results
    uses e.g. '1'+bye (full/allocated bye), '0'+'not paired' (zero-point/absent),
    and forfeits like '+ - -' (white wins) / '- - +' (black wins)."""
    s = str(raw).strip().replace("½", "0.5").lower()
    compact = s.replace(" ", "")
    if not has_black:
        if compact in ("1", "+"):
            return "U"          # pairing-allocated / full-point bye (1.0)
        if compact in ("0.5", "="):
            return "H"          # half-point bye
        if compact in ("0", "-", ""):
            return "Z"          # zero-point bye (absent / not paired)
        return "U"
    table = {
        "1-0": "1-0", "0-1": "0-1", "0.5-0.5": "0.5-0.5",
        "+-": "1-0F", "+--": "1-0F", "1-0f": "1-0F",
        "-+": "0-1F", "--+": "0-1F", "0-1f": "0-1F",
        "--": "0-0", "0-0": "0-0",
    }
    if compact in table:
        return table[compact]
    raise ValueError(f"unrecognized result '{raw}'")


def parse_chess_results_roster(rows: list[list]) -> list[dict]:
    """Parse a Chess-Results 'starting rank list' export (roster only) — columns
    like: No. | Name | FideID | FED | Rtg | sex | Typ | Club/City. Metadata rows
    above the header are skipped."""
    hdr_i = None
    for i, r in enumerate(rows[:40]):
        low = [_norm(c) for c in r if c is not None]
        if "name" in low and any(k in low for k in ("no", "startno", "rtg", "rating", "seed")):
            hdr_i = i
            break
    if hdr_i is None:
        raise ValueError("could not find the player-list header row (No./Name/Rtg)")
    return parse_players(rows[hdr_i:])


def parse_chess_results(rows: list[list]) -> tuple[list[dict], list[dict]]:
    # locate header row + column indices
    hdr_i = None
    for i, r in enumerate(rows[:40]):
        cells = [str(c).strip().lower() if c is not None else "" for c in r]
        if "white" in cells and "black" in cells:
            hdr_i = i
            break
    if hdr_i is None:
        raise ValueError("could not find the 'White/Black' header row")
    hdr = [str(c).strip().lower() if c is not None else "" for c in rows[hdr_i]]

    def idx(pred, start=0):
        for j in range(start, len(hdr)):
            if pred(hdr[j]):
                return j
        return None

    c_board = idx(lambda c: c.startswith("bo")) or 0
    c_white = idx(lambda c: c == "white")
    c_black = idx(lambda c: c == "black")
    c_res = idx(lambda c: c.startswith("result"))
    c_wrtg = idx(lambda c: c == "rtg", c_white + 1)
    c_brtg = idx(lambda c: c == "rtg", c_black + 1)
    if None in (c_white, c_black, c_res):
        raise ValueError("Chess-Results header missing White/Black/Result columns")

    ratings: dict[str, int] = {}
    raw_games: list[tuple] = []      # (round, board, white, black_or_None, result)
    cur_round = None
    BYE_WORDS = {"", "bye", "spielfrei", "not paired", "-1", "none"}

    for r in rows:
        c0 = str(r[0]).strip() if r and r[0] is not None else ""
        m = c0.lower()
        if m.startswith("round") and any(ch.isdigit() for ch in c0):
            cur_round = int("".join(ch for ch in c0 if ch.isdigit()))
            continue
        # skip repeated header rows (each round block has its own)
        low = [str(c).strip().lower() if c is not None else "" for c in r]
        if "white" in low and "black" in low:
            continue
        # a game row: board is an int and there's a white name
        if not (isinstance(r[0], (int, float)) and c_white < len(r) and r[c_white]):
            continue
        if cur_round is None:
            cur_round = 1
        board = int(r[0])
        white = str(r[c_white]).strip()
        black_raw = str(r[c_black]).strip() if c_black < len(r) and r[c_black] is not None else ""
        black = None if black_raw.lower() in BYE_WORDS else black_raw
        res = str(r[c_res]).strip() if c_res < len(r) and r[c_res] is not None else ""
        wr = _to_int(r[c_wrtg]) if c_wrtg and c_wrtg < len(r) else None
        br = _to_int(r[c_brtg]) if c_brtg and c_brtg < len(r) else None
        # always register both players (rating may legitimately be 0)
        ratings[white] = max(ratings.get(white, 0), wr or 0)
        if black:
            ratings[black] = max(ratings.get(black, 0), br or 0)
        raw_games.append((cur_round, board, white, black, res))

    if not raw_games:
        raise ValueError("no game rows found in the Chess-Results export")

    # seed players by rating desc, then name (matches Chess-Results starting rank)
    names = sorted(ratings.keys(), key=lambda n: (-(ratings.get(n) or 0), n))
    seed = {n: i + 1 for i, n in enumerate(names)}
    players = [{"start_no": seed[n], "name": n, "rating": ratings.get(n) or 0,
                "fide_id": None, "sex": None, "title": None, "federation": None}
               for n in names]

    pairings = []
    for (rnd, board, white, black, res) in raw_games:
        pairings.append({
            "round": rnd, "board": board,
            "white": seed[white], "black": (seed[black] if black else None),
            "result": _cr_result(res, has_black=black is not None),
        })
    return players, pairings


def parse_xlsx(data: bytes) -> tuple[list[dict], list[dict]]:
    wb = load_workbook(io.BytesIO(data), data_only=True)

    # Chess-Results.com single-sheet export?
    for name in wb.sheetnames:
        rows = _rows_from_xlsx_sheet(wb[name])
        if _is_chess_results(rows):
            has_wb = any(
                "white" in [str(c).strip().lower() if c is not None else "" for c in r]
                and "black" in [str(c).strip().lower() if c is not None else "" for c in r]
                for r in rows[:40]
            )
            if has_wb:
                return parse_chess_results(rows)          # pairings/results export
            return parse_chess_results_roster(rows), []   # starting-rank list (roster only)

    # Single-sheet roster (starting-rank list, no games): a header row with a
    # Name column and a seed/rating column, but no round/white/black columns.
    # Chess-Results exports this without the site banner, so detect it by shape.
    if len(wb.sheetnames) == 1:
        rows = _rows_from_xlsx_sheet(wb[wb.sheetnames[0]])
        for r in rows[:40]:
            low = [_norm(c) for c in r if c is not None]
            if not low:
                continue
            if ("name" in low
                    and any(k in low for k in ("no", "startno", "seed", "rtg", "rating"))
                    and not any(k in low for k in ("round", "white", "black"))):
                return parse_chess_results_roster(rows), []
            break  # first non-empty row wasn't a roster header → not this shape

    sheets = {name.lower(): name for name in wb.sheetnames}

    def find(*keys):
        for k in keys:
            for low, real in sheets.items():
                if k in low:
                    return wb[real]
        return None

    ws_players = find("player", "roster", "seed")
    ws_pairings = find("pairing", "result", "game", "round")
    if ws_players is None or ws_pairings is None:
        # fall back to first two sheets in order
        names = wb.sheetnames
        ws_players = ws_players or wb[names[0]]
        ws_pairings = ws_pairings or (wb[names[1]] if len(names) > 1 else wb[names[0]])
    return (parse_players(_rows_from_xlsx_sheet(ws_players)),
            parse_pairings(_rows_from_xlsx_sheet(ws_pairings)))


def parse_csvs(players_text: str, pairings_text: str) -> tuple[list[dict], list[dict]]:
    return (parse_players(_rows_from_csv(players_text)),
            parse_pairings(_rows_from_csv(pairings_text)))


# --------------------------- record (dict) parsing ------------------------- #
def _pick(d: dict, field: str, names: set):
    for k, v in d.items():
        nk = _norm(k)
        if nk == field.replace("_", "") or nk in names:
            return v
    return None


def parse_players_records(records: list[dict]) -> list[dict]:
    out = []
    for d in records:
        sno = _to_int(_pick(d, "start_no", _PLAYER_ALIASES["start_no"]))
        name = _pick(d, "name", _PLAYER_ALIASES["name"])
        if sno is None or not name:
            continue
        out.append({
            "start_no": sno, "name": str(name).strip(),
            "rating": _to_int(_pick(d, "rating", _PLAYER_ALIASES["rating"])) or 0,
            "fide_id": (str(_pick(d, "fide_id", _PLAYER_ALIASES["fide_id"])).strip()
                        if _pick(d, "fide_id", _PLAYER_ALIASES["fide_id"]) not in (None, "") else None),
            "sex": (str(_pick(d, "sex", _PLAYER_ALIASES["sex"])).strip().lower()[:1]
                    if _pick(d, "sex", _PLAYER_ALIASES["sex"]) else None),
            "title": (str(_pick(d, "title", _PLAYER_ALIASES["title"])).strip()
                      if _pick(d, "title", _PLAYER_ALIASES["title"]) else None),
            "federation": (str(_pick(d, "federation", _PLAYER_ALIASES["federation"])).strip()
                           if _pick(d, "federation", _PLAYER_ALIASES["federation"]) else None),
        })
    if not out:
        raise ValueError("no player records parsed (need start_no + name)")
    return out


def parse_pairings_records(records: list[dict]) -> list[dict]:
    out = []
    for d in records:
        rnd = _to_int(_pick(d, "round", _PAIRING_ALIASES["round"]))
        white = _to_int(_pick(d, "white", _PAIRING_ALIASES["white"]))
        if rnd is None or white is None:
            continue
        black = _to_int(_pick(d, "black", _PAIRING_ALIASES["black"]))
        if black == 0:
            black = None
        result = _pick(d, "result", _PAIRING_ALIASES["result"])
        board = _to_int(_pick(d, "board", _PAIRING_ALIASES["board"]))
        out.append({"round": rnd, "board": board, "white": white, "black": black,
                    "result": str(result).strip() if result is not None else ""})
    return out


def parse_json(data: bytes) -> tuple[list[dict], list[dict]]:
    """Accept either a list of players, or an object with players/pairings arrays.
    Keys are matched leniently, e.g. {"players":[...], "pairings":[...]} or
    {"roster":[...], "results":[...]}."""
    obj = json.loads(data.decode("utf-8-sig"))
    if isinstance(obj, list):
        return parse_players_records(obj), []
    if not isinstance(obj, dict):
        raise ValueError("JSON must be a list of players or an object")
    players_raw = (obj.get("players") or obj.get("roster") or obj.get("Players") or [])
    pairings_raw = (obj.get("pairings") or obj.get("results") or obj.get("games")
                    or obj.get("Pairings") or [])
    return parse_players_records(players_raw), parse_pairings_records(pairings_raw)
